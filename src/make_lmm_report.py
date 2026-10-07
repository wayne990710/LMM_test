"""LMM 分析報告（Word + PDF）：研究變因 → 資料與樣本數 → 模型 → 結果 → 解讀與限制。

表格數字直接讀 results/LMM/ 的結果表。
執行：python src/make_lmm_report.py（需先跑過 lmm_main.py；轉 PDF 需要 Microsoft Word）
輸出：results/LMM/LMM 分析結果.docx、.pdf
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import data as D
from make_word import ACCENT, FONT_EN, INK, MUTED, _apply_fonts, add_rich, keep_full_resolution, spacing

R = D.ROOT / "results" / "LMM"
OUT = R / "LMM 分析結果.docx"
FIGS = [R / "01_各節次樣本數.png", R / "02_LMM 係數.png"]
W = 21.0 - 2 * 2.0   # A4 直式，左右邊界 2 cm
LINE = 17


# ---------------------------------------------------------------- 排版
def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT_EN
    st.font.size = Pt(11)
    _apply_fonts(st.element.rPr.rFonts)


def h1(doc, text, new_page=False):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.keep_with_next = True
    spacing(p, before=0 if new_page else 12, after=4)
    add_rich(p, text, size=14, bold_all=True, color=ACCENT)


def para(doc, text, size=11, color=INK, indent=False, after=4):
    p = doc.add_paragraph()
    spacing(p, after=after, exact=LINE)
    if indent:
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
    add_rich(p, text, size=size, color=color)
    return p


def bullets(doc, items, size=11):
    for t in items:
        para(doc, "・" + t, size=size, indent=True, after=2)


def _shade(cell, hex_):
    tc = cell._element.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), hex_)
    tc.append(s)


def table(doc, rows, widths, header_fill="E6F2F0", size=9.5, bold_rows=()):
    doc.paragraphs[-1].paragraph_format.keep_with_next = True   # 表名與表格放在同一頁
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            p = c.paragraphs[0]
            spacing(p, after=0, exact=13)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_rich(p, str(v), size=size, bold_all=(i == 0 or i in bold_rows))
            if i == 0:
                _shade(c, header_fill)
            if i < len(rows) - 1:          # 表格不跨頁：每列都與下一列放在同一頁
                p.paragraph_format.keep_with_next = True
        trp = t.rows[i]._tr.get_or_add_trPr()
        trp.append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p, before=4, after=2, line=1.0)
    p.paragraph_format.keep_with_next = True   # 圖與圖說放在同一頁
    p.add_run().add_picture(str(path), width=Cm(W - 1.0))
    para(doc, caption, size=9.5, color=MUTED, after=8)


def fmt_p(p):
    return "< .001" if p < 0.001 else f"{p:.3f}".replace("0.", ".", 1)


def fmt(v, unit):
    d = 3 if unit in ("比例", "log ms") else 2 if unit == "分（1–7）" else 1
    return f"{v:+.{d}f}"


# ---------------------------------------------------------------- 內容
def build():
    size = pd.read_csv(R / "lmm_sample_size.csv")
    coef = pd.read_csv(R / "lmm_coefficients.csv")
    fit = pd.read_csv(R / "lmm_model_fit.csv").set_index("依變量")
    vif = pd.read_csv(R / "lmm_vif.csv").set_index("依變量")
    sens = pd.read_csv(R / "lmm_sensitivity_no_sleep.csv")
    sw = pd.read_csv(R / "ac_switch_time.csv", index_col=0).iloc[:, 0]

    doc = Document()
    setup(doc)
    p = doc.add_paragraph()
    spacing(p, after=2)
    add_rich(p, "教室環境對學生疲勞、認知表現與心率的影響", size=18, bold_all=True)
    p = doc.add_paragraph()
    spacing(p, after=8)
    add_rich(p, "線性混合效應模型（LMM）分析結果｜資料期間 2026-09-17 ～ 10-02｜12 位學生、14 個施測節次", size=11,
             color=MUTED)

    # 1. 研究變因
    h1(doc, "一、研究變因")
    table(doc, [
        ["類別", "變項", "測量方式"],
        ["自變量", "二氧化碳（CO₂）", "教室兩臺 SCD30 感測器；係數以每 100 ppm 表示"],
        ["", "溫度", "同上，°C"],
        ["", "濕度", "同上，相對濕度 %"],
        ["依變量", "疲勞程度", "自覺疲勞量表「我現在的疲勞程度」（1–7 分）"],
        ["", "反應時間", "Stroop 答對題的平均反應時間（只用 200–3000 ms 的題目）"],
        ["", "正確率", "Stroop 答對題數 ÷ 總題數"],
        ["", "干擾分數", "Stroop 不一致題 − 中性題的平均反應時間"],
        ["", "瞬時心率", "手環或心電貼片的心率，每 5 分鐘平均"],
        ["", "心率變異度", "心電貼片逐拍 RR 間隔算出的 RMSSD，每 5 分鐘一筆，取對數"],
        ["控制變因", "人數", "研究代碼對照表的現場人數"],
        ["", "冷氣", "研究代碼對照表（09-24 起）；09-21 ～ 09-23 依 CO₂ 上升速度推定"],
        ["", "睡眠時長", "疲勞量表的睡眠題（不到 6 小時、6～8 小時、8 小時以上，換算為 5、7、8.5 小時）"],
        ["", "睡眠品質", "疲勞量表的睡眠品質題"],
    ], [2.2, 3.0, 11.8])
    para(doc, "環境值（CO₂、溫度、濕度）的對應時間：疲勞取填答前 3 分鐘、Stroop 取施測當下、"
              "心率與心率變異度取同一個 5 分鐘。", size=10, color=MUTED)

    # 2. 資料與樣本數
    h1(doc, "二、資料與樣本數")
    para(doc, "本研究的 Stroop 測驗與疲勞量表都在**下課後施測**。有時下課時間比較趕，部分學生沒有測到，"
              "所以每一節課的人數不同，各依變量的樣本數也不一樣。LMM 以學生為隨機效應，可以處理這種"
              "「每個人測到的次數不一樣」的資料，所以**所有有資料的學生與節次都納入分析**，沒有只挑測驗完整的人。")
    para(doc, "各依變量的樣本數如下。「納入模型」需要自變量與控制變因都有值，所以比「有測到」少：")
    rows = [["依變量", "有測到", "納入模型", "學生數", "主要缺漏原因"]]
    why = {"疲勞程度": "09-17 沒有 CO₂ 資料",
           "反應時間": "沒填疲勞量表就沒有睡眠資料；09-17 沒有 CO₂",
           "正確率": "同上", "干擾分數": "同上",
           "瞬時心率": "沒填量表的節次缺睡眠；09-17 沒有 CO₂",
           "心率變異度 log(RMSSD)": "只有心電貼片能算；缺睡眠"}
    for _, r in size.iterrows():
        unit = "筆（5 分鐘）" if r["資料層級"] == "5 分鐘" else "人次"
        rows.append([r["依變量"].replace(" log(RMSSD)", ""), f"{r['有測到的筆數']} {unit}",
                     f"{r['納入模型的筆數']} {unit}", f"{r['納入模型的學生']} 人", why[r["依變量"]]])
    table(doc, rows, [2.6, 2.9, 2.9, 1.6, 7.0])
    bullets(doc, [
        "疲勞量表排除注意力檢核題未通過的問卷（7 份，皆為同一位學生），所以疲勞與睡眠只有 11 人。",
        "心率與心率變異度要知道「哪個裝置是哪位學生戴的」。研究代碼對照表從 09-24 起才寫裝置編號；"
        "09-21 ～ 09-23 只記了字母代號，當時沒有記錄對應的編號，所以這段期間的生理資料無法連到學生。",
        "心率變異度只能用心電貼片的逐拍資料計算（手環只輸出平滑過的心率），且只用乾淨心跳 ≥ 80% 的 5 分鐘時段，"
        "所以只有 5 位學生。",
    ], size=10.5)
    figure(doc, FIGS[0], "圖 1　各節次有資料的學生數。顏色越深代表人數越多；「—」為該節沒有資料。")

    # 3. 模型
    h1(doc, "三、LMM 模型")
    para(doc, "六個依變量各建一個模型，三個自變量與四個控制變因同時放入：")
    p = para(doc, "依變量 ＝ CO₂ ＋ 溫度 ＋ 濕度 ＋ 人數 ＋ 冷氣 ＋ 睡眠時長 ＋ 睡眠品質 ＋（學生隨機截距）", size=11)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bullets(doc, [
        "學生隨機截距：每位學生有自己的基準（例如有人本來反應就比較快），模型比較的是同一位學生在不同環境下的變化。",
        "心率與心率變異度是每 5 分鐘一筆，同一節課內的多筆資料彼此相關，所以再加上「學生 × 節次」的隨機截距。",
        f"冷氣：「後半冷氣」的節次，Stroop 與量表都在下課後施測，算開冷氣；心率與心率變異度則依 CO₂ 開始累積的時間"
        f"區分前後（推定開啟時間：09-21 上午 {pd.Timestamp(sw['09-21 AM']):%H:%M}、"
        f"09-29 上午 {pd.Timestamp(sw['09-29 AM']):%H:%M}）。",
        "CO₂ 來源：兩臺感測器在該節都有完整資料就用平均，否則用有資料的那一臺（每節固定一個來源）。",
        "估計方法：REML（statsmodels MixedLM）。圖 2 的標準化係數 ＝ 係數 × 預測變項標準差 ÷ 依變量標準差，"
        "讓不同單位的變項可以互相比較。",
    ], size=10.5)

    # 4. 結果
    h1(doc, "四、結果", new_page=True)
    para(doc, "表 2　三個自變量的係數（已控制其他兩個自變量與四個控制變因）。括號內為 95% 信賴區間，**粗體**為 p < .05。")
    ivs = ["CO₂（每 100 ppm）", "溫度（°C）", "濕度（%）"]
    rows = [["依變量（單位）", "CO₂（每 +100 ppm）", "溫度（每 +1 °C）", "濕度（每 +1 %）"]]
    for dv in dict.fromkeys(coef["依變量"]):
        g = coef[coef["依變量"] == dv].set_index("變項")
        unit = g["單位"].iloc[0]
        cells = []
        for v in ivs:
            r = g.loc[v]
            s = f"{fmt(r['係數'], unit)}\n({fmt(r['95% CI 下限'], unit)}, {fmt(r['95% CI 上限'], unit)})\np = {fmt_p(r.p)}"
            cells.append(f"**{s}**" if r.p < 0.05 else s)
        rows.append([f"{dv.replace(' log(RMSSD)', '')}\n（{unit}）"] + cells)
    t = table(doc, rows, [4.0, 4.3, 4.3, 4.3], size=9.5)
    for row in t.rows[1:]:   # 粗體標記處理：add_rich 只認 **…** 在同一段，換行拆段後重設
        for c in row.cells:
            txt = c.text
            if txt.startswith("**"):
                c.paragraphs[0].clear()
                add_rich(c.paragraphs[0], txt.strip("*"), size=9.5, bold_all=True, color=INK)
    para(doc, "心率變異度為 log(RMSSD)：係數 × 100 約等於 RMSSD 變化的百分比（例如 +0.18 ≈ +20%）。", size=9.5,
         color=MUTED)

    para(doc, "表 3　控制變因的係數（p < .05 者）：")
    sig = coef[(coef["角色"] == "控制變因") & (coef.p < 0.05)]
    rows = [["依變量", "控制變因", "係數（95% CI）", "p"]]
    for _, r in sig.iterrows():
        rows.append([r["依變量"].replace(" log(RMSSD)", ""), r["變項"],
                     f"{fmt(r['係數'], r['單位'])}（{fmt(r['95% CI 下限'], r['單位'])}, {fmt(r['95% CI 上限'], r['單位'])}）",
                     fmt_p(r.p)])
    table(doc, rows, [4.0, 4.0, 6.0, 2.0])

    figure(doc, FIGS[1], "圖 2　六個 LMM 的標準化係數。上方三列（彩色）為自變量，下方四列（灰色）為控制變因；"
                         "線為 95% 信賴區間，實心點為 p < .05。")

    para(doc, "表 4　模型解釋力。邊際 R² 只看固定效應（環境與控制變因），條件 R² 再加上學生（與節次）之間的差異：")
    rows = [["依變量", "筆數", "學生", "邊際 R²", "條件 R²"]]
    for dv, r in fit.iterrows():
        rows.append([dv.replace(" log(RMSSD)", ""), int(r["筆數"]), int(r["學生"]), f"{r['邊際 R²']:.2f}",
                     f"{r['條件 R²']:.2f}"])
    table(doc, rows, [4.5, 2.5, 2.5, 3.0, 3.0])

    # 5. 解讀
    h1(doc, "五、結果摘要")
    bullets(doc, [
        "**CO₂**：在同時控制溫度、濕度與四個控制變因後，CO₂ 對六個依變量都沒有達到顯著。",
        "**溫度**：溫度較高時，干擾分數較小（每 +1 °C −67 ms）、瞬時心率較低（−4.0 bpm）、心率變異度較高（約 +20%）。",
        "**濕度**：濕度較高時，瞬時心率較高（每 +1 % +0.6 bpm）；對其他依變量沒有顯著影響。",
        "**疲勞程度**：只有睡眠時長顯著（每多睡 1 小時，疲勞 −0.61 分）；環境三個變項都不顯著。",
        "**控制變因**：開冷氣時瞬時心率較低（−4.0 bpm）。",
        "**反應時間與正確率**：沒有任何變項達到顯著；學生之間的差異很大（反應時間的條件 R² 0.62，邊際 R² 只有 0.04）。",
    ])

    # 6. 限制
    h1(doc, "六、解讀時要注意")
    v_co2 = vif["CO₂（每 100 ppm）"].max()
    bullets(doc, [
        f"**自變量之間彼此相關**：開冷氣時會關窗，CO₂ 上升、溫度與濕度同時下降。三者同時放入模型後，"
        f"共同變動的部分很難分給哪一個變項，所以每個係數的信賴區間都比較寬。變異數膨脹因子（VIF）"
        f"CO₂ 最高 {v_co2:.1f}、所有變項最高 {vif.max().max():.1f}，在可接受範圍（< 10），但仍會削弱檢定力。",
        "**人數幾乎沒有變化**：有 CO₂ 資料的節次中，現場人數幾乎都是 23 人，只有 09-29 上午是 22 人，"
        "所以人數的係數其實只反映那一節課，不代表人數的效果。",
        "**冷氣在生理資料中幾乎都開著**：心率與心率變異度只有 09-24 起的資料，這段期間都有開冷氣，只有 09-29 上午"
        "開冷氣前的時段是關的。冷氣對心率的係數主要來自那一節課，要保守解讀。",
        "**心率變異度只有 5 位學生**，結果較不穩定。",
    ])
    para(doc, "敏感度分析：睡眠資料來自疲勞量表，沒填量表的 Stroop 與心率資料會被排除。拿掉兩個睡眠控制變因、"
              "納入所有有資料的筆數後，三個自變量的結果如下（p < .05 以粗體標示）：")
    rows = [["依變量", "筆數／學生", "CO₂", "溫度", "濕度"]]
    unit = coef.groupby("依變量")["單位"].first()
    for dv in dict.fromkeys(sens["依變量"]):
        g = sens[sens["依變量"] == dv].set_index("變項")
        cells = []
        for v in ivs:
            r = g.loc[v]
            s = f"{fmt(r['係數'], unit[dv])}（p = {fmt_p(r.p)}）"
            cells.append(f"**{s}**" if r.p < 0.05 else s)
        rows.append([dv.replace(" log(RMSSD)", ""), f"{int(g['筆數'].iloc[0])}／{int(g['學生'].iloc[0])}"] + cells)
    table(doc, rows, [3.6, 2.4, 3.7, 3.7, 3.6])
    bullets(doc, [
        "納入更多資料後，反應時間的 CO₂（每 +100 ppm +4.5 ms）與溫度（每 +1 °C +54 ms）、心率變異度的 CO₂"
        "（約 −3.6%）變成顯著；瞬時心率與疲勞程度的結論不變。",
        "結果會隨納入的樣本改變，代表目前的資料量還不足以穩定地區分 CO₂、溫度與濕度各自的影響。",
    ], size=10.5)

    keep_full_resolution(doc)
    doc.save(OUT)
    print(OUT)


def to_pdf():
    pdf = OUT.with_suffix(".pdf")
    ps = ("$w = New-Object -ComObject Word.Application; $w.Visible = $false; $w.DisplayAlerts = 0; "
          f"try {{ $d = $w.Documents.Open('{OUT}', $false, $true); $d.ExportAsFixedFormat('{pdf}', 17, $false, 0); "
          "$d.ComputeStatistics(2); $d.Close($false) } finally { $w.Quit() }")
    out = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, check=True)
    # Word 匯出會把圖降到約 150 dpi：依出現順序換回 300 dpi 原圖
    import pymupdf
    doc = pymupdf.open(pdf)
    imgs = [(page, im[0]) for page in doc for im in page.get_images(full=True)]
    assert len(imgs) == len(FIGS), f"PDF 中有 {len(imgs)} 張圖"
    for (page, xref), f in zip(imgs, FIGS):
        page.replace_image(xref, filename=str(f))
    tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False).name
    doc.save(tmp, garbage=4, deflate=True)
    doc.close()
    shutil.move(tmp, pdf)
    print(pdf, out.stdout.strip().splitlines()[-1], "頁")


if __name__ == "__main__":
    build()
    to_pdf()
